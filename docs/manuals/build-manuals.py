"""
Build the three Fshikta Dental user manuals into print-ready HTML, PDF and DOCX.

    pip install python-docx markdown
    python docs/manuals/build-manuals.py

Outputs land in docs/manuals/build/. PDF generation uses headless Chrome or Edge,
whichever is found; if neither is installed the HTML is still produced and can be
printed to PDF from any browser.
"""

from __future__ import annotations

import base64
import html as html_mod
import mimetypes
import re
import shutil
import subprocess
import sys
import tempfile
import time
from pathlib import Path

import markdown
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
BUILD = HERE / "build"
IMAGES = HERE / "images"

MANUALS = [
    ("01-Receptionist-Manual.md", "Receptionist Manual"),
    ("02-Doctors-and-Nurses-Manual.md", "Doctors and Nurses Manual"),
    ("03-Manager-and-Accountants-Manual.md", "Manager and Accountants Manual"),
]

CHROME_CANDIDATES = [
    r"C:\Program Files\Google\Chrome\Application\chrome.exe",
    r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
    r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
]

CSS = """
@page { size: A4; margin: 18mm 16mm 20mm; }
* { box-sizing: border-box; }
body {
  font: 11pt/1.55 "Segoe UI", Calibri, Arial, sans-serif;
  color: #1a2230; margin: 0; padding: 0 4mm;
}
h1 { font-size: 24pt; color: #0f5c7a; margin: 0 0 4pt; page-break-after: avoid; }
h2 { font-size: 15pt; color: #0f5c7a; margin: 22pt 0 6pt;
     padding-bottom: 3pt; border-bottom: 1.5pt solid #bfe0eb; page-break-after: avoid; }
h3 { font-size: 12.5pt; color: #214055; margin: 14pt 0 4pt; page-break-after: avoid; }
h4 { font-size: 11pt; color: #214055; margin: 12pt 0 3pt; }
p, li { orphans: 2; widows: 2; }
a { color: #0f5c7a; text-decoration: none; }
code { background: #eef3f6; padding: 1pt 3pt; border-radius: 3px;
       font: 9.5pt "Consolas", monospace; }
pre { background: #f4f7f9; border: 0.75pt solid #d6e2e8; border-left: 3pt solid #0f5c7a;
      padding: 8pt 10pt; border-radius: 4px; font: 9pt/1.4 "Consolas", monospace;
      white-space: pre; overflow-x: auto; page-break-inside: avoid; }
table { border-collapse: collapse; width: 100%; margin: 8pt 0 12pt;
        font-size: 10pt; page-break-inside: avoid; }
th { background: #0f5c7a; color: #fff; text-align: left; font-weight: 600;
     padding: 5pt 7pt; border: 0.5pt solid #0f5c7a; }
td { padding: 4.5pt 7pt; border: 0.5pt solid #cfdde4; vertical-align: top; }
tbody tr:nth-child(even) td { background: #f6fafb; }
blockquote { margin: 10pt 0; padding: 7pt 12pt; background: #fff8e6;
             border-left: 3pt solid #e0a800; border-radius: 3px; page-break-inside: avoid; }
blockquote p { margin: 0; }
img { max-width: 100%; height: auto; border: 0.75pt solid #cfdde4;
      border-radius: 4px; margin: 6pt 0 2pt; display: block; page-break-inside: avoid; }
figure { margin: 10pt 0 14pt; page-break-inside: avoid; }
figcaption { font-size: 9pt; color: #5b6b7a; font-style: italic; margin-top: 2pt; }
hr { border: 0; border-top: 0.75pt solid #d6e2e8; margin: 16pt 0; }
ul, ol { padding-left: 20pt; }
li { margin: 2pt 0; }
.missing-shot { background: #fdecea; border: 0.75pt dashed #d33; border-radius: 4px;
                color: #a3231b; padding: 10pt; font-size: 9.5pt; }
"""


def md_to_html(md_text: str, title: str, embed_images: bool) -> str:
    body = markdown.markdown(
        md_text,
        extensions=["tables", "fenced_code", "toc", "sane_lists", "attr_list"],
    )
    body = wrap_figures(body, embed_images)
    return (
        "<!doctype html>\n<html lang=\"en\"><head><meta charset=\"utf-8\">"
        f"<title>{html_mod.escape(title)}</title><style>{CSS}</style></head>"
        f"<body>{body}</body></html>"
    )


def wrap_figures(body: str, embed_images: bool) -> str:
    """Turn <img alt="Figure n - caption"> into a <figure> with a caption, and
    replace missing screenshots with a visible placeholder so gaps are obvious."""

    def repl(match: re.Match[str]) -> str:
        alt = match.group("alt")
        src = match.group("src")
        path = (HERE / src).resolve()
        caption = html_mod.escape(alt)
        if not path.exists():
            return (
                f'<div class="missing-shot"><strong>Screenshot missing:</strong> '
                f"{caption} <br><code>{html_mod.escape(src)}</code></div>"
            )
        if embed_images:
            mime = mimetypes.guess_type(path.name)[0] or "image/png"
            data = base64.b64encode(path.read_bytes()).decode("ascii")
            src = f"data:{mime};base64,{data}"
        return f'<figure><img src="{src}" alt="{caption}"><figcaption>{caption}</figcaption></figure>'

    return re.sub(
        r'<img alt="(?P<alt>[^"]*)" src="(?P<src>[^"]+)"\s*/?>', repl, body
    )


def find_chrome() -> str | None:
    for candidate in CHROME_CANDIDATES:
        if Path(candidate).exists():
            return candidate
    for name in ("chrome", "msedge", "chromium"):
        found = shutil.which(name)
        if found:
            return found
    return None


def html_to_pdf(html_path: Path, pdf_path: Path, chrome: str) -> bool:
    with tempfile.TemporaryDirectory() as profile:
        cmd = [
            chrome,
            "--headless=new",
            "--disable-gpu",
            "--no-sandbox",
            f"--user-data-dir={profile}",
            "--no-pdf-header-footer",
            "--virtual-time-budget=20000",
            f"--print-to-pdf={pdf_path}",
            html_path.as_uri(),
        ]
        result = subprocess.run(cmd, capture_output=True, text=True, timeout=180)
    for _ in range(10):
        if pdf_path.exists() and pdf_path.stat().st_size > 0:
            return True
        time.sleep(0.5)
    sys.stderr.write(f"  PDF failed: {result.stderr.strip()[:300]}\n")
    return False


# ---------------------------------------------------------------- DOCX writer

TABLE_ROW_RE = re.compile(r"^\s*\|(.+)\|\s*$")
SEP_ROW_RE = re.compile(r"^\s*\|[\s:|-]+\|\s*$")
IMG_RE = re.compile(r"^!\[(?P<alt>[^\]]*)\]\((?P<src>[^)]+)\)\s*$")
INLINE_RE = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`|\[.+?\]\(.+?\))")


def add_inline(paragraph, text: str) -> None:
    """Render **bold**, *italic*, `code` and [links](...) into one paragraph."""
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        elif part.startswith("[") and "](" in part:
            label = part[1 : part.index("](")]
            run = paragraph.add_run(label)
            run.font.color.rgb = RGBColor(0x0F, 0x5C, 0x7A)
        else:
            paragraph.add_run(part)


def flush_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=cols)
    table.style = "Light Grid Accent 1"
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for col in range(cols):
            text = row[col] if col < len(row) else ""
            cell = cells[col]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            if index == 0:
                run = paragraph.add_run(re.sub(r"\*\*", "", text))
                run.bold = True
            else:
                add_inline(paragraph, text)
            for run in paragraph.runs:
                run.font.size = Pt(9.5)
    doc.add_paragraph()


def md_to_docx(md_text: str, title: str, out_path: Path) -> None:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Segoe UI"
    normal.font.size = Pt(10.5)

    lines = md_text.splitlines()
    table_rows: list[list[str]] = []
    in_code = False
    code_buffer: list[str] = []

    def flush_code() -> None:
        if not code_buffer:
            return
        paragraph = doc.add_paragraph()
        run = paragraph.add_run("\n".join(code_buffer))
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        code_buffer.clear()

    for raw in lines:
        line = raw.rstrip()

        if line.startswith("```"):
            in_code = not in_code
            if not in_code:
                flush_code()
            continue
        if in_code:
            code_buffer.append(raw)
            continue

        if TABLE_ROW_RE.match(line):
            if SEP_ROW_RE.match(line):
                continue
            table_rows.append([c.strip() for c in line.strip().strip("|").split("|")])
            continue
        if table_rows:
            flush_table(doc, table_rows)
            table_rows = []

        if not line.strip():
            continue

        image = IMG_RE.match(line.strip())
        if image:
            src = (HERE / image.group("src")).resolve()
            caption = image.group("alt")
            if src.exists():
                doc.add_picture(str(src), width=Inches(6.2))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = cap.add_run(caption)
                run.italic = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x5B, 0x6B, 0x7A)
            else:
                placeholder = doc.add_paragraph()
                run = placeholder.add_run(f"[Screenshot missing: {caption}]")
                run.italic = True
                run.font.color.rgb = RGBColor(0xA3, 0x23, 0x1B)
            continue

        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            doc.add_heading(line[level:].strip(), level=min(level, 4))
            continue
        if line.strip() in {"---", "***", "___"}:
            continue
        if line.startswith(">"):
            paragraph = doc.add_paragraph(style="Intense Quote")
            add_inline(paragraph, line.lstrip("> ").strip())
            continue
        stripped = line.strip()
        if re.match(r"^[-*] \[[ xX]\] ", stripped):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline(paragraph, re.sub(r"^[-*] \[[ xX]\] ", "", stripped))
            continue
        if stripped.startswith(("- ", "* ")):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline(paragraph, stripped[2:])
            continue
        if re.match(r"^\d+[.)] ", stripped):
            paragraph = doc.add_paragraph(style="List Number")
            add_inline(paragraph, re.sub(r"^\d+[.)] ", "", stripped))
            continue

        paragraph = doc.add_paragraph()
        add_inline(paragraph, stripped)

    if table_rows:
        flush_table(doc, table_rows)
    flush_code()
    doc.save(out_path)


def main() -> int:
    BUILD.mkdir(parents=True, exist_ok=True)
    chrome = find_chrome()
    if not chrome:
        print("No Chrome or Edge found — HTML and DOCX only, no PDF.")

    missing_total = 0
    for filename, title in MANUALS:
        source = HERE / filename
        if not source.exists():
            print(f"skipped (not found): {filename}")
            continue
        md_text = source.read_text(encoding="utf-8")
        stem = source.stem

        missing = [
            m.group("src")
            for m in re.finditer(r"!\[[^\]]*\]\((?P<src>images/[^)]+)\)", md_text)
            if not (HERE / m.group("src")).exists()
        ]
        missing_total += len(missing)

        html_path = BUILD / f"{stem}.html"
        html_path.write_text(md_to_html(md_text, title, embed_images=False), encoding="utf-8")

        pdf_source = BUILD / f".{stem}.embedded.html"
        pdf_source.write_text(md_to_html(md_text, title, embed_images=True), encoding="utf-8")

        docx_path = BUILD / f"{stem}.docx"
        md_to_docx(md_text, title, docx_path)

        made_pdf = False
        if chrome:
            made_pdf = html_to_pdf(pdf_source, BUILD / f"{stem}.pdf", chrome)
        pdf_source.unlink(missing_ok=True)

        note = f" ({len(missing)} screenshots missing)" if missing else ""
        print(f"{title}: html + docx{' + pdf' if made_pdf else ''}{note}")

    # The HTML build needs the screenshots beside it.
    if IMAGES.exists():
        target = BUILD / "images"
        target.mkdir(exist_ok=True)
        for image in IMAGES.glob("*.png"):
            shutil.copy2(image, target / image.name)

    print(f"\noutput: {BUILD}")
    if missing_total:
        print(
            f"{missing_total} screenshot(s) not yet captured — run "
            "docs/manuals/capture-screenshots.mjs with the app running."
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
